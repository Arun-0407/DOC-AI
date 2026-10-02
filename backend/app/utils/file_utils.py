import os
import aiofiles
from fastapi import UploadFile
import zipfile

async def save_upload_file(file: UploadFile, upload_dir: str) -> str:
    os.makedirs(upload_dir, exist_ok=True)
    file_path = os.path.join(upload_dir, file.filename)
    async with aiofiles.open(file_path, 'wb') as out_file:
        content = await file.read()
        await out_file.write(content)
    return file_path

async def extract_zip(file: UploadFile, upload_dir: str) -> list:
    os.makedirs(upload_dir, exist_ok=True)
    zip_path = os.path.join(upload_dir, file.filename)
    
    async with aiofiles.open(zip_path, 'wb') as out_file:
        content = await file.read()
        await out_file.write(content)
        
    extracted_files = []
    with zipfile.ZipFile(zip_path, 'r') as zip_ref:
        for info in zip_ref.infolist():
            if not info.is_dir():
                with zip_ref.open(info) as f:
                    content = f.read().decode('utf-8', errors='ignore')
                    extracted_files.append({"filename": info.filename, "content": content})
                    
    return extracted_files

def detect_abap_type(content: str, filename: str = '') -> str:
    """Detect the type of SAP/source object from content and filename."""
    ext = ''
    if filename:
        import os
        ext = os.path.splitext(filename)[1].lower()

    # Detect by file extension first (most reliable)
    if ext == '.pdf':
        return 'pdf_document'
    if ext in ('.doc', '.docx'):
        return 'word_document'
    if ext in ('.xls', '.xlsx'):
        return 'excel_document'
    if ext == '.csv':
        return 'csv_document'
    if ext == '.json':
        return 'json_document'
    if ext in ('.md', '.markdown'):
        return 'markdown_document'
    if ext in ('.sql',):
        return 'sql_script'
    if ext in ('.html', '.htm'):
        return 'html_document'

    content_upper = content.upper()
    # XML-based SAP objects
    if content.strip().startswith('<?xml') or content.strip().startswith('<'):
        if 'SMARTFORMS' in content_upper or 'SMART_FORM' in content_upper:
            return 'smart_form'
        elif 'FUGR' in content_upper or 'FUNCTION GROUP' in content_upper:
            return 'function_group'
        elif 'CLAS' in content_upper or 'CLASS' in content_upper:
            return 'class'
        elif 'PROG' in content_upper or 'REPORT' in content_upper:
            return 'abap_report'
        elif 'TABL' in content_upper or 'DBTAB' in content_upper:
            return 'table'
        elif 'DDLS' in content_upper or 'CDS' in content_upper:
            return 'cds_view'
        elif 'WDYN' in content_upper or 'WEB DYNPRO' in content_upper:
            return 'web_dynpro'
        elif 'ENHO' in content_upper or 'ENHANCEMENT' in content_upper:
            return 'enhancement'
        else:
            return 'xml_object'
    # Plain ABAP source
    if 'REPORT ' in content_upper or 'PROGRAM ' in content_upper:
        return 'abap_report'
    elif 'CLASS ' in content_upper and 'DEFINITION' in content_upper:
        return 'class'
    elif 'FUNCTION-POOL' in content_upper:
        return 'function_group'
    elif 'FUNCTION ' in content_upper and 'ENDFUNCTION' in content_upper:
        return 'function_module'
    elif 'INTERFACE ' in content_upper:
        return 'interface'
    elif 'METHOD ' in content_upper:
        return 'method'
    return 'abap_report'

def read_file_content(filepath: str) -> str:
    """Read file content. For binary files (PDF, Word, Excel) extracts text.
    For text files reads directly. Never raises — always returns something useful."""
    import os
    ext = os.path.splitext(filepath)[1].lower()

    # PDF text extraction
    if ext == '.pdf':
        try:
            import pdfplumber
            text_parts = []
            with pdfplumber.open(filepath) as pdf:
                for page in pdf.pages:
                    t = page.extract_text()
                    if t:
                        text_parts.append(t)
            return '\n'.join(text_parts) if text_parts else _read_as_text(filepath)
        except Exception:
            return _read_as_text(filepath)

    # Word .docx text extraction
    if ext == '.docx':
        try:
            from docx import Document as DocxDoc
            doc = DocxDoc(filepath)
            return '\n'.join(p.text for p in doc.paragraphs if p.text.strip())
        except Exception:
            return _read_as_text(filepath)

    # Excel .xlsx text extraction
    if ext in ('.xlsx', '.xls'):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
            lines = []
            for sheet in wb.worksheets:
                lines.append(f'[Sheet: {sheet.title}]')
                for row in sheet.iter_rows(values_only=True):
                    row_str = '\t'.join(str(c) if c is not None else '' for c in row)
                    if row_str.strip():
                        lines.append(row_str)
            wb.close()
            return '\n'.join(lines)
        except Exception:
            return _read_as_text(filepath)

    # All other files — read as text
    return _read_as_text(filepath)


def _read_as_text(filepath: str) -> str:
    """Fallback: read file as UTF-8 text, ignoring encoding errors."""
    try:
        with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read()
    except Exception:
        # Last resort: read as binary and decode
        try:
            with open(filepath, 'rb') as f:
                return f.read().decode('utf-8', errors='ignore')
        except Exception:
            return f'[Could not read file: {filepath}]'
