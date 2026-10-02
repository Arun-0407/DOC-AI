from docx import Document as DocxDocument
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import io
import json
from app.models.document import Document

SAP_BLUE = RGBColor(0x0F, 0x4C, 0x81)
SAP_DARK = RGBColor(0x0A, 0x16, 0x28)


def _add_section_heading(doc: DocxDocument, number: int, title: str):
    """Add a styled section heading."""
    heading = doc.add_heading(f"{number}. {title}", level=1)
    for run in heading.runs:
        run.font.color.rgb = SAP_BLUE


def _add_table(doc: DocxDocument, headers: list, rows: list):
    """Add a formatted table."""
    if not rows:
        doc.add_paragraph("No data available.", style='List Bullet')
        return
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Header row
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(9)
    # Data rows
    for row_idx, row_data in enumerate(rows):
        for col_idx, value in enumerate(row_data):
            table.rows[row_idx + 1].cells[col_idx].text = str(value or "")


def _safe_get(data, key, default=""):
    """Safely get a value from dict or return default."""
    if data is None:
        return default
    if isinstance(data, dict):
        return data.get(key, default)
    return default


def generate_docx(document: Document) -> bytes:
    doc = DocxDocument()

    # ── Cover Page ──
    doc.add_paragraph("")
    doc.add_paragraph("")
    title = doc.add_heading(document.title or 'SAP Technical Documentation', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.color.rgb = SAP_BLUE

    cover = document.cover_data or {}
    if cover:
        subtitle = doc.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = subtitle.add_run(f"Program: {_safe_get(cover, 'program_name', 'N/A')}")
        run.font.size = Pt(14)
        run.font.color.rgb = SAP_BLUE

        meta_items = [
            ("Module", _safe_get(cover, "module")),
            ("Object Type", _safe_get(cover, "object_type")),
            ("Author", _safe_get(cover, "author")),
            ("Version", _safe_get(cover, "version")),
            ("Created", _safe_get(cover, "creation_date")),
        ]
        for label, value in meta_items:
            if value:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(f"{label}: {value}")
                run.font.size = Pt(10)

    doc.add_page_break()

    # ── 1. Business View ──
    bv = document.business_view or {}
    _add_section_heading(doc, 1, "Business View")
    if bv:
        doc.add_heading("Objective", level=2)
        doc.add_paragraph(_safe_get(bv, "objective"))
        doc.add_heading("Executive Summary", level=2)
        doc.add_paragraph(_safe_get(bv, "executive_summary"))
        doc.add_heading("Business Purpose", level=2)
        doc.add_paragraph(_safe_get(bv, "business_purpose"))
        if bv.get("business_rules"):
            doc.add_heading("Business Rules", level=2)
            for rule in bv["business_rules"]:
                doc.add_paragraph(str(rule), style='List Bullet')
        if bv.get("key_benefits"):
            doc.add_heading("Key Benefits", level=2)
            for benefit in bv["key_benefits"]:
                doc.add_paragraph(str(benefit), style='List Bullet')
    doc.add_page_break()

    # ── 2. Technical View ──
    tv = document.technical_view or {}
    _add_section_heading(doc, 2, "Technical View")
    if tv:
        meta = _safe_get(tv, "program_metadata", {})
        if meta:
            doc.add_heading("Program Metadata", level=2)
            for k, v in meta.items() if isinstance(meta, dict) else []:
                doc.add_paragraph(f"{k}: {v}", style='List Bullet')

        sel = _safe_get(tv, "selection_screen", {})
        if sel and isinstance(sel, dict):
            params = sel.get("parameters", [])
            if params:
                doc.add_heading("Selection Screen Parameters", level=2)
                rows = [[p.get("name", ""), p.get("type", ""), str(p.get("obligatory", ""))] for p in params if isinstance(p, dict)]
                _add_table(doc, ["Name", "Type", "Obligatory"], rows)
    doc.add_page_break()

    # ── 3. Source Code Analysis ──
    sa = document.source_analysis or {}
    _add_section_heading(doc, 3, "Source Code Analysis")
    routines = _safe_get(sa, "routines", [])
    if routines and isinstance(routines, list):
        rows = [[r.get("name", ""), r.get("type", ""), r.get("complexity", ""), r.get("purpose", "")] for r in routines if isinstance(r, dict)]
        _add_table(doc, ["Name", "Type", "Complexity", "Purpose"], rows)
    doc.add_page_break()

    # ── 4. Database Analysis ──
    da = document.database_analysis or {}
    _add_section_heading(doc, 4, "Database Analysis")
    tables = _safe_get(da, "tables", [])
    if tables and isinstance(tables, list):
        rows = []
        for t in tables:
            if isinstance(t, dict):
                crud = t.get("crud_usage", {})
                crud_str = ", ".join([k.upper() for k, v in crud.items() if v]) if isinstance(crud, dict) else ""
                rows.append([t.get("name", ""), t.get("description", ""), crud_str])
        _add_table(doc, ["Table", "Description", "CRUD Operations"], rows)
    doc.add_page_break()

    # ── 5. API & Dependencies ──
    api_dep = document.api_dependency or {}
    _add_section_heading(doc, 5, "API & Dependencies")
    for category in ["function_modules", "bapis", "rfcs", "odata_services", "rest_apis"]:
        items = _safe_get(api_dep, category, [])
        if items and isinstance(items, list):
            doc.add_heading(category.replace("_", " ").title(), level=2)
            rows = [[i.get("name", ""), i.get("description", "")] for i in items if isinstance(i, dict)]
            _add_table(doc, ["Name", "Description"], rows)
    doc.add_page_break()

    # ── 6. Process Flow ──
    pf = document.process_flow or {}
    _add_section_heading(doc, 6, "Process Flow")
    doc.add_paragraph("Process flow diagrams are available in the HTML/PDF export and online viewer.")
    doc.add_page_break()

    # ── 7. Integration View ──
    iv = document.integration_view or {}
    _add_section_heading(doc, 7, "Integration View")
    for category in ["sap_to_sap", "sap_to_non_sap", "rfc", "bapi", "idoc", "rest", "odata"]:
        items = _safe_get(iv, category, [])
        if items and isinstance(items, list):
            doc.add_heading(category.replace("_", " ").title(), level=2)
            rows = [[i.get("name", ""), i.get("type", ""), i.get("direction", ""), i.get("description", "")] for i in items if isinstance(i, dict)]
            _add_table(doc, ["Name", "Type", "Direction", "Description"], rows)
    doc.add_page_break()

    # ── 8. Security Review ──
    sr = document.security_review or {}
    _add_section_heading(doc, 8, "Security Review")
    score = _safe_get(sr, "security_score", "N/A")
    doc.add_paragraph(f"Security Score: {score}/100")
    for category in ["auth_objects", "sensitive_data_access", "rfc_risks", "hardcoded_secrets", "compliance_issues"]:
        items = _safe_get(sr, category, [])
        if items and isinstance(items, list):
            doc.add_heading(category.replace("_", " ").title(), level=2)
            for item in items:
                if isinstance(item, dict):
                    doc.add_paragraph(f"{item.get('name', '')}: {item.get('description', '')}", style='List Bullet')
    doc.add_page_break()

    # ── 9. Performance Review ──
    pr = document.performance_review or {}
    _add_section_heading(doc, 9, "Performance Review")
    pscore = _safe_get(pr, "performance_score", "N/A")
    doc.add_paragraph(f"Performance Score: {pscore}/100")
    for category in ["select_in_loop", "nested_loops", "optimization_opportunities"]:
        items = _safe_get(pr, category, [])
        if items and isinstance(items, list):
            doc.add_heading(category.replace("_", " ").title(), level=2)
            rows = [[i.get("location", ""), i.get("issue", ""), i.get("suggestion", "")] for i in items if isinstance(i, dict)]
            _add_table(doc, ["Location", "Issue", "Suggestion"], rows)
    doc.add_page_break()

    # ── 10. Reference Data ──
    rd = document.reference_data or {}
    _add_section_heading(doc, 10, "Reference Data")
    for category in ["transactions", "tables", "function_modules", "classes", "interfaces", "message_classes"]:
        items = _safe_get(rd, category, [])
        if items and isinstance(items, list):
            doc.add_heading(category.replace("_", " ").title(), level=2)
            rows = [[i.get("name", ""), i.get("description", "")] for i in items if isinstance(i, dict)]
            _add_table(doc, ["Name", "Description"], rows)
    doc.add_page_break()

    # ── 11. Modification History ──
    mh = document.modification_history or {}
    _add_section_heading(doc, 11, "Modification History")
    history = _safe_get(mh, "history", [])
    if history and isinstance(history, list):
        rows = [[h.get("date", ""), h.get("object", ""), h.get("modification", ""), h.get("author", ""), h.get("impact", "")] for h in history if isinstance(h, dict)]
        _add_table(doc, ["Date", "Object", "Modification", "Author", "Impact"], rows)
    doc.add_page_break()

    # ── 12. AI Recommendations ──
    ar = document.ai_recommendations or {}
    _add_section_heading(doc, 12, "AI Recommendations")
    scores = ["code_quality_score", "maintainability_score", "performance_score",
              "security_score", "s4hana_readiness_score", "rap_migration_score"]
    for s in scores:
        val = _safe_get(ar, s, "N/A")
        doc.add_paragraph(f"{s.replace('_', ' ').title()}: {val}/100", style='List Bullet')

    suggestions = _safe_get(ar, "refactoring_suggestions", [])
    if suggestions and isinstance(suggestions, list):
        doc.add_heading("Refactoring Suggestions", level=2)
        for sug in suggestions:
            if isinstance(sug, dict):
                doc.add_paragraph(f"[{sug.get('priority', '')}] {sug.get('title', '')}: {sug.get('description', '')}", style='List Bullet')

    clean_recs = _safe_get(ar, "clean_abap_recommendations", [])
    if clean_recs and isinstance(clean_recs, list):
        doc.add_heading("Clean ABAP Recommendations", level=2)
        for rec in clean_recs:
            if isinstance(rec, dict):
                doc.add_paragraph(f"{rec.get('title', '')}: {rec.get('description', '')}", style='List Bullet')

    # Save to bytes
    f = io.BytesIO()
    doc.save(f)
    return f.getvalue()
